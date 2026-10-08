import io
import json
import os
import re
import streamlit as st
import pypdf
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai
from google.genai import types

# ---------------------------------------------------------
# Konfigurasi Halaman Streamlit
# ---------------------------------------------------------
st.set_page_config(
    page_title="Generator Silabus & RPP Kitab",
    page_icon="📖",
    layout="wide"
)

st.title("📖 Generator Silabus & RPP Otomatis dari Kitab")
st.caption("Unggah dokumen kitab (PDF/Docx), teks manual, atau file Silabus .docx untuk merancang Perangkat Ajar Bertahap!")

# ---------------------------------------------------------
# Inisialisasi Google GenAI Client
# ---------------------------------------------------------
api_key = os.environ.get("GEMINI_API_KEY")
if not api_key and "GEMINI_API_KEY" in st.secrets:
    api_key = st.secrets["GEMINI_API_KEY"]

if not api_key:
    st.error("⚠️ GEMINI_API_KEY tidak ditemukan di Secrets atau Environment Variables.")
    st.stop()

client = genai.Client(api_key=api_key)

# ---------------------------------------------------------
# Fungsi Helper: Membaca Teks PDF Berdasarkan Rentang Halaman
# ---------------------------------------------------------
def parse_page_range(range_str: str, total_pages: int):
    if not range_str.strip():
        return set(range(total_pages))
    
    selected_pages = set()
    parts = range_str.split(',')
    for part in parts:
        part = part.strip()
        if '-' in part:
            sub = part.split('-')
            if len(sub) == 2 and sub[0].isdigit() and sub[1].isdigit():
                start = max(1, int(sub[0]))
                end = min(total_pages, int(sub[1]))
                for p in range(start, end + 1):
                    selected_pages.add(p - 1)
        elif part.isdigit():
            p = int(part)
            if 1 <= p <= total_pages:
                selected_pages.add(p - 1)
    
    return sorted(list(selected_pages))

def extract_text_from_pdf(pdf_file, page_range_str: str = "") -> str:
    reader = pypdf.PdfReader(pdf_file)
    total_pages = len(reader.pages)
    
    page_indices = parse_page_range(page_range_str, total_pages)
    if not page_indices:
        page_indices = list(range(total_pages))
        
    text = ""
    for idx in page_indices:
        extracted = reader.pages[idx].extract_text()
        if extracted:
            text += f"\n--- Halaman {idx + 1} ---\n" + extracted
    return text

# ---------------------------------------------------------
# Fungsi Helper: Ekstraksi Teks dari File Word (.docx) Umum
# ---------------------------------------------------------
def extract_text_from_docx(docx_file) -> str:
    doc = Document(docx_file)
    full_text = []
    for p in doc.paragraphs:
        if p.text.strip():
            full_text.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text)

# ---------------------------------------------------------
# Fungsi Helper: Parse Silabus dari File .docx yang Diunggah
# ---------------------------------------------------------
def parse_silabus_from_docx(docx_file) -> dict:
    doc = Document(docx_file)
    silabus_list = []
    
    if doc.tables:
        table = doc.tables[0]
        headers = [cell.text.strip().lower() for cell in table.rows[0].cells]
        
        for row in table.rows[1:]:
            row_data = [cell.text.strip() for cell in row.cells]
            if len(row_data) >= len(headers):
                item = {
                    "ptm": int(row_data[0]) if row_data[0].isdigit() else len(silabus_list) + 1,
                    "bab": row_data[1] if len(row_data) > 1 else "",
                    "halaman": row_data[2] if len(row_data) > 2 else "",
                    "alokasi_waktu": row_data[3] if len(row_data) > 3 else "",
                    "metode": row_data[4] if len(row_data) > 4 else "",
                    "capaian": row_data[5] if len(row_data) > 5 else "",
                    "indikator": row_data[6] if len(row_data) > 6 else "",
                    "evaluasi": row_data[7] if len(row_data) > 7 else ""
                }
                silabus_list.append(item)
                
    if not silabus_list:
        current_item = {}
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if "pertemuan" in text.lower() or "ptm" in text.lower():
                if current_item:
                    silabus_list.append(current_item)
                current_item = {
                    "ptm": len(silabus_list) + 1,
                    "bab": text,
                    "halaman": "-",
                    "alokasi_waktu": "2 x 45 Menit",
                    "metode": "Bandongan / Sorogan / Diskusi",
                    "capaian": text,
                    "indikator": "Memahami isi materi",
                    "evaluasi": "Tes lisan"
                }
        if current_item:
            silabus_list.append(current_item)

    if not silabus_list:
        silabus_list = [{
            "ptm": 1,
            "bab": "Materi Pembelajaran dari Dokumen",
            "halaman": "1-end",
            "alokasi_waktu": "2 x 45 Menit",
            "metode": "Klasik",
            "capaian": "Pemahaman awal",
            "indikator": "Keaktifan santri",
            "evaluasi": "Tanya jawab"
        }]

    return {
        "nama_kitab": "Kitab dari Dokumen Unggahan",
        "fan_ilmu": "Umum / Diniyah",
        "tingkat_kelas": "Umum",
        "total_pertemuan": str(len(silabus_list)),
        "silabus": silabus_list
    }

# ---------------------------------------------------------
# Fungsi Helper: Membuat Dokumen Word (.docx) untuk Silabus
# ---------------------------------------------------------
def create_silabus_docx(data: dict) -> io.BytesIO:
    doc = Document()
    
    p_lembaga = doc.add_paragraph()
    p_lembaga.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_lembaga.add_run("THE RUFIDZ INDONESIA\n")
    run_inst.bold = True
    run_inst.font.size = Pt(14)
    run_inst.font.name = 'Arial'
    
    run_title = p_lembaga.add_run("SILABUS PEMBELAJARAN KITAB")
    run_title.bold = True
    run_title.font.size = Pt(12)
    run_title.font.name = 'Arial'

    meta_text = (
        f"Nama Kitab    : {data.get('nama_kitab', '-')}\n"
        f"Fan Ilmu      : {data.get('fan_ilmu', '-')}\n"
        f"Kelas / Tahap : {data.get('tingkat_kelas', '-')}\n"
        f"Total Ptm     : {data.get('total_pertemuan', '-')}"
    )
    p_meta = doc.add_paragraph(meta_text)
    p_meta.style.font.name = 'Arial'
    p_meta.style.font.size = Pt(11)
    doc.add_paragraph("\n")

    silabus_items = data.get("silabus", [])
    if silabus_items:
        table = doc.add_table(rows=1, cols=8)
        table.style = 'Table Grid'
        headers = ["Ptm", "Bab / Fasal", "Halaman", "Alokasi Waktu", "Metode", "Capaian", "Indikator", "Evaluasi"]
        
        hdr_cells = table.rows[0].cells
        for i, h_text in enumerate(headers):
            hdr_cells[i].text = h_text
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9)
            hdr_cells[i].paragraphs[0].runs[0].font.name = 'Arial'

        for item in silabus_items:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.get("ptm", ""))
            row_cells[1].text = str(item.get("bab", ""))
            row_cells[2].text = str(item.get("halaman", ""))
            row_cells[3].text = str(item.get("alokasi_waktu", ""))
            row_cells[4].text = str(item.get("metode", ""))
            row_cells[5].text = str(item.get("capaian", ""))
            row_cells[6].text = str(item.get("indikator", ""))
            row_cells[7].text = str(item.get("evaluasi", ""))
            
            for cell in row_cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        run.font.name = 'Arial'

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ---------------------------------------------------------
# Fungsi Helper: Membuat File Excel (.xlsx) untuk Jurnal Harian KBM
# ---------------------------------------------------------
def create_jurnal_excel(data: dict) -> io.BytesIO:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        fan_text = data.get('fan_ilmu', 'SHARAF').upper()
        kitab_text = data.get('nama_kitab', '')
        
        rows_data = []
        rows_data.append([f"JURNAL PEMBELAJARAN {fan_text} THE RUFIDZ INDONESIA", "", "", "", "", "", ""])
        rows_data.append([f"SEMESTER GANJIL T.A 2026/2027 (Kitab: {kitab_text})", "", "", "", "", "", ""])
        rows_data.append(["", "", "", "", "", "", ""])
        
        rows_data.append([
            "MATERI POKOK", 
            "ALOKASI WAKTU", 
            "Tanggal", 
            "NAMA PENGAJAR/PENGGANTI", 
            "CATATAN MATERI\n(Diisi jika diperlukan)", 
            "TTD", 
            "Keterangan"
        ])
        
        silabus_items = data.get("silabus", [])
        for item in silabus_items:
            rows_data.append([
                str(item.get("bab", "")),
                str(item.get("alokasi_waktu", "")),
                "......./....../2026",
                "",
                "",
                "",
                ""
            ])
            
        df_jurnal = pd.DataFrame(rows_data)
        df_jurnal.to_excel(writer, sheet_name='Jurnal KBM', index=False, header=False)
        
    output.seek(0)
    return output

# ---------------------------------------------------------
# Fungsi Helper: Format Dokumen Word (.docx) untuk RPP Tunggal
# ---------------------------------------------------------
def create_single_rpp_docx(data: dict, rpp: dict) -> io.BytesIO:
    doc = Document()
    
    p_lembaga = doc.add_paragraph()
    p_lembaga.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_lembaga.add_run("THE RUFIDZ INDONESIA\n")
    run_inst.bold = True
    run_inst.font.size = Pt(14)
    run_inst.font.name = 'Arial'
    
    run_rpp_title = p_lembaga.add_run("RENCANA PELAKSANAAN PEMBELAJARAN (RPP)")
    run_rpp_title.bold = True
    run_rpp_title.font.size = Pt(12)
    run_rpp_title.font.name = 'Arial'

    rpp_id = (
        f"Sekolah       : {rpp.get('sekolah', 'MDT / Pesantren Rufidz Tahfidz & Diniyah Indonesia')}\n"
        f"Mata Pelajaran: {data.get('fan_ilmu', '')} ({data.get('nama_kitab', '')})\n"
        f"Kelas / Tahap : {data.get('tingkat_kelas', '')}\n"
        f"Materi Pokok  : {rpp.get('materi_pokok', '')}\n"
        f"Alokasi Waktu : {rpp.get('alokasi_waktu', '')}"
    )
    p_ident = doc.add_paragraph(rpp_id)
    p_ident.style.font.name = 'Arial'
    p_ident.style.font.size = Pt(11)

    p_tujuan_head = doc.add_paragraph()
    r_tujuan = p_tujuan_head.add_run("A. Tujuan Pembelajaran")
    r_tujuan.bold = True
    r_tujuan.font.size = Pt(11)
    r_tujuan.font.name = 'Arial'
    
    for t in rpp.get("tujuan_pembelajaran", []):
        p = doc.add_paragraph(f"• {t}")
        p.style.font.name = 'Arial'
        p.style.font.size = Pt(11)

    p_langkah_head = doc.add_paragraph()
    r_langkah = p_langkah_head.add_run("B. Langkah-Langkah Pembelajaran")
    r_langkah.bold = True
    r_langkah.font.size = Pt(11)
    r_langkah.font.name = 'Arial'
    
    p_pend = doc.add_paragraph()
    r_pend = p_pend.add_run("1. Kegiatan Pendahuluan")
    r_pend.bold = True
    r_pend.font.size = Pt(11)
    r_pend.font.name = 'Arial'
    
    for p_item in rpp.get("langkah_pembelajaran", {}).get("pendahuluan", []):
        p = doc.add_paragraph(f"• {p_item}")
        p.style.font.name = 'Arial'
        p.style.font.size = Pt(11)

    p_inti_head = doc.add_paragraph()
    r_inti_head = p_inti_head.add_run("2. Kegiatan Inti")
    r_inti_head.bold = True
    r_inti_head.font.size = Pt(11)
    r_inti_head.font.name = 'Arial'
    
    inti_list = rpp.get("langkah_pembelajaran", {}).get("inti", [])
    if inti_list:
        table_inti = doc.add_table(rows=1, cols=2)
        table_inti.style = 'Table Grid'
        hdr = table_inti.rows[0].cells
        hdr[0].text = "Aspek"
        hdr[1].text = "Kegiatan Pembelajaran"
        hdr[0].paragraphs[0].runs[0].font.bold = True
        hdr[1].paragraphs[0].runs[0].font.bold = True
        hdr[0].paragraphs[0].runs[0].font.size = Pt(10)
        hdr[1].paragraphs[0].runs[0].font.size = Pt(10)
        
        for item in inti_list:
            r_cells = table_inti.add_row().cells
            r_cells[0].text = item.get("aspek", "")
            r_cells[1].text = item.get("kegiatan", "")
            r_cells[0].paragraphs[0].runs[0].font.size = Pt(10)
            r_cells[1].paragraphs[0].runs[0].font.size = Pt(10)

    p_penut = doc.add_paragraph()
    r_penut = p_penut.add_run("\n3. Kegiatan Penutup")
    r_penut.bold = True
    r_penut.font.size = Pt(11)
    r_penut.font.name = 'Arial'
    
    for p_item in rpp.get("langkah_pembelajaran", {}).get("penutup", []):
        p = doc.add_paragraph(f"• {p_item}")
        p.style.font.name = 'Arial'
        p.style.font.size = Pt(11)

    p_penilaian_head = doc.add_paragraph()
    r_penilaian = p_penilaian_head.add_run("C. Penilaian Hasil Pembelajaran")
    r_penilaian.bold = True
    r_penilaian.font.size = Pt(11)
    r_penilaian.font.name = 'Arial'
    
    penilaian = rpp.get("penilaian", {})
    st_p1 = doc.add_paragraph(f"• Penilaian Sikap: {penilaian.get('sikap', '')}")
    st_p1.style.font.name = 'Arial'
    st_p1.style.font.size = Pt(11)
    
    st_p2 = doc.add_paragraph(f"• Penilaian Pengetahuan: {penilaian.get('pengetahuan', '')}")
    st_p2.style.font.name = 'Arial'
    st_p2.style.font.size = Pt(11)
    
    st_p3 = doc.add_paragraph(f"• Penilaian Keterampilan: {penilaian.get('keterampilan', '')}")
    st_p3.style.font.name = 'Arial'
    st_p3.style.font.size = Pt(11)

    doc.add_paragraph("\n")
    ttd_table = doc.add_table(rows=2, cols=2)
    ttd_cells_0 = ttd_table.rows[0].cells
    ttd_cells_1 = ttd_table.rows[1].cells
    
    ttd_cells_0[0].text = "Mengetahui,\nKepala Sekolah"
    ttd_cells_0[1].text = "\nPengampu Mapel"
    ttd_cells_1[0].text = "\n\n( Mudir Sekolah )"
    ttd_cells_1[1].text = "\n\n( Ustadz Pengampu )"
    
    for row in ttd_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(11)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ---------------------------------------------------------
# Form Input
# ---------------------------------------------------------
col1, col2 = st.columns([1, 1])

with col1:
    nama_kitab = st.text_input("Nama Kitab / Pelajaran:", placeholder="Contoh: Al-Qawaidus Sharfiyyah")
    fan_ilmu = st.text_input("Fan Ilmu:", placeholder="Contoh: Sharaf")
    tingkat_kelas = st.text_input("Tingkat / Kelas:", placeholder="Contoh: Kelas 7")
    
    col_ptm, col_waktu = st.columns(2)
    with col_ptm:
        total_pertemuan_input = st.text_input("Total Pertemuan:", placeholder="Contoh: 134 Pertemuan")
    with col_waktu:
        alokasi_waktu_input = st.text_input("Alokasi Waktu per Pertemuan:", placeholder="Contoh: 2 x 45 Menit")

with col2:
    metode_input = st.radio(
        "Pilih Sumber Masukan:",
        ["📁 Upload Dokumen Kitab (PDF/Docx)", "✍️ Ketik / Paste Teks", "📄 Upload Silabus .docx"],
        horizontal=True
    )
    
    uploaded_doc_kitab = None
    kitab_text_input = ""
    page_range_input = ""
    uploaded_docx_silabus = None
    
    if metode_input == "📁 Upload Dokumen Kitab (PDF/Docx)":
        uploaded_doc_kitab = st.file_uploader("Unggah Dokumen Kitab (PDF atau Docx, Max 10MB):", type=["pdf", "docx"])
        page_range_input = st.text_input(
            "📄 Rentang Halaman yang Ingin Dianalisis (Opsional untuk PDF):",
            placeholder="Contoh: 1-30 (Kosongkan jika semua halaman)"
        )
    elif metode_input == "✍️ Ketik / Paste Teks":
        kitab_text_input = st.text_area(
            "Masukkan / Paste Teks Kitab di Sini:",
            placeholder="Salin dan tempelkan teks materi/kitab di sini...",
            height=200
        )
    else:
        uploaded_docx_silabus = st.file_uploader("Unggah File Silabus Berformat .docx:", type=["docx"])

if "rpp_cache" not in st.session_state:
    st.session_state["rpp_cache"] = {}

# ---------------------------------------------------------
# Tombol Aksi Pembuatan Silabus / Load Silabus Word
# ---------------------------------------------------------
if metode_input == "📄 Upload Silabus .docx":
    if st.button("📂 1. Load Data Silabus dari Dokumen .docx", type="primary"):
        if not uploaded_docx_silabus:
            st.error("Harap unggah file Silabus berformat .docx terlebih dahulu.")
        else:
            try:
                parsed_data = parse_silabus_from_docx(uploaded_docx_silabus)
                parsed_data["nama_kitab"] = nama_kitab if nama_kitab else parsed_data["nama_kitab"]
                parsed_data["fan_ilmu"] = fan_ilmu if fan_ilmu else parsed_data["fan_ilmu"]
                parsed_data["tingkat_kelas"] = tingkat_kelas if tingkat_kelas else parsed_data["tingkat_kelas"]
                
                st.session_state["result_data"] = parsed_data
                st.session_state["rpp_cache"] = {}
                st.success("✅ Silabus berhasil dimuat dari file Word! Silakan cek tab di bawah.")
            except Exception as docx_err:
                st.error(f"Gagal membaca file .docx: {docx_err}")
else:
    btn_label = "📂 1. Buat Silabus dari Dokumen" if metode_input == "📁 Upload Dokumen Kitab (PDF/Docx)" else "🚀 1. Buat Silabus Pembelajaran Terlebih Dahulu"
    if st.button(btn_label, type="primary"):
        final_text = ""
        if metode_input == "📁 Upload Dokumen Kitab (PDF/Docx)":
            if not uploaded_doc_kitab:
                st.error("Harap unggah berkas dokumen kitab (PDF atau Docx) terlebih dahulu.")
                st.stop()
            with st.spinner("Membaca dan menganalisis dokumen kitab..."):
                file_extension = uploaded_doc_kitab.name.split(".")[-1].lower()
                if file_extension == "pdf":
                    final_text = extract_text_from_pdf(uploaded_doc_kitab, page_range_input)
                elif file_extension == "docx":
                    final_text = extract_text_from_docx(uploaded_doc_kitab)
                
                if not final_text.strip():
                    st.error("Teks pada dokumen tidak terbaca.")
                    st.stop()
        else:
            if not kitab_text_input.strip():
                st.error("Harap masukkan teks kitab terlebih dahulu.")
                st.stop()
            final_text = kitab_text_input

        st.session_state["pdf_text"] = final_text
        st.session_state["nama_kitab"] = nama_kitab
        st.session_state["fan_ilmu"] = fan_ilmu
        st.session_state["tingkat_kelas"] = tingkat_kelas
        st.session_state["total_pertemuan"] = total_pertemuan_input
        st.session_state["alokasi_waktu"] = alokasi_waktu_input
        st.session_state["rpp_cache"] = {}

        with st.spinner("Gemini AI sedang menyusun Silabus dari seluruh teks kitab..."):
            prompt = f"""
            Anda adalah seorang pakar kurikulum madrasah/pesantren.
            Buatkan pemetaan Silabus secara lengkap dan proporsional untuk seluruh pertemuan dari teks kitab di bawah ini.
            
            --- DETAIL INPUT ---
            - Nama Kitab: {nama_kitab}
            - Fan Ilmu: {fan_ilmu}
            - Tingkat / Kelas: {tingkat_kelas}
            - Total Pertemuan: {total_pertemuan_input if total_pertemuan_input else 'Sesuaikan'}
            - Alokasi Waktu per Pertemuan: {alokasi_waktu_input if alokasi_waktu_input else '2 x 45 Menit'}
            --- TEKS KITAB ---
            {final_text}
            --------------------
            
            Keluarkan respons HANYA dalam bentuk JSON valid dengan struktur berikut tanpa teks lain:
            {{
                "nama_kitab": "{nama_kitab}",
                "fan_ilmu": "{fan_ilmu}",
                "tingkat_kelas": "{tingkat_kelas}",
                "total_pertemuan": "{total_pertemuan_input}",
                "silabus": [
                    {{
                        "ptm": 1,
                        "bab": "Nama bab/fasal dari teks",
                        "halaman": "Rentang halaman",
                        "alokasi_waktu": "{alokasi_waktu_input}",
                        "metode": "Metode pembelajaran",
                        "capaian": "Capaian indikator",
                        "indikator": "Indikator ketercapaian",
                        "evaluasi": "Bentuk evaluasi"
                    }}
                ]
            }}
            """

            import time
            max_retries = 3
            retry_delay = 2
            success = False
            response = None

            for attempt in range(max_retries):
                try:
                    response = client.models.generate_content(
                        model="gemini-3.6-flash",
                        contents=prompt,
                        config=types.GenerateContentConfig(
                            max_output_tokens=65536, temperature=0.7
                        ),
                    )
                    success = True
                    break
                except Exception as e:
                    if "503" in str(e) and attempt < max_retries - 1:
                        time.sleep(retry_delay)
                        retry_delay *= 2
                        continue
                    else:
                        st.error(f"Gagal memproses AI / JSON: {str(e)}")
                        break

            if success and response:
                try:
                    raw_text = response.text.strip()
                    if "```json" in raw_text:
                        raw_text = raw_text.split("```json")[1].split("```")[0].strip()
                    elif "```" in raw_text:
                        raw_text = raw_text.split("```")[1].split("```")[0].strip()

                    start_idx = raw_text.find("{")
                    end_idx = raw_text.rfind("}")
                    if start_idx != -1 and end_idx != -1:
                        raw_text = raw_text[start_idx : end_idx + 1]

                    data = json.loads(raw_text)
                    st.session_state["result_data"] = data
                    st.success("✅ Silabus berhasil dibuat dari teks utuh! Silakan lihat tab di bawah.")
                except Exception as json_err:
                    st.error(f"Gagal melakukan parsing data JSON dari AI: {json_err}")
                    with st.expander("🔍 Lihat Mentah Respons AI (for debugging)"):
                        st.text(response.text)

# ---------------------------------------------------------
# Tampilan Hasil & Pilihan Tab
# ---------------------------------------------------------
if "result_data" in st.session_state:
    data = st.session_state["result_data"]
    
    st.markdown("---")
    st.header(f"📖 Kurikulum: {data.get('nama_kitab', '')}")
    
    tab_silabus, tab_jurnal, tab_rpp = st.tabs([
        "📑 Silabus Pembelajaran", 
        "📊 Jurnal Pembelajaran KBM (Excel)", 
        "📝 Generator RPP Per Pertemuan"
    ])
    
    with tab_silabus:
        silabus_list = data.get("silabus", [])
        if silabus_list:
            df_silabus = pd.DataFrame(silabus_list)
            column_mapping = {
                "ptm": "Ptm", "bab": "Bab / Fasal", "halaman": "Halaman",
                "alokasi_waktu": "Alokasi Waktu", "metode": "Metode Klasik",
                "capaian": "Capaian Indikator", "indikator": "Indikator Ketercapaian", "evaluasi": "Bentuk Evaluasi"
            }
            df_silabus.rename(columns={k: v for k, v in column_mapping.items() if k in df_silabus.columns}, inplace=True)
            df_silabus = df_silabus.loc[:, ~df_silabus.columns.duplicated()]
            st.dataframe(df_silabus, use_container_width=True, hide_index=True)
            
            # --- FITUR: REVISI SILABUS DENGAN PROMPT USER ---
            st.markdown("---")
            st.markdown("### 🛠️ Perbaiki atau Sesuaikan Silabus dengan Perintah AI")
            st.caption("Ketik perintah khusus jika Anda ingin merevisi silabus di atas (contoh: *'Pecah pertemuan 2 menjadi dua sesi khusus praktik tashrif'*).")
            
            correction_prompt = st.text_area(
                "Instruksi / Perintah Perbaikan Silabus:",
                placeholder="Tulis instruksi revisi silabus di sini...",
                key="silabus_correction_input"
            )
            
            if st.button("✨ Terapkan Perbaikan Silabus", type="primary"):
                if not correction_prompt.strip():
                    st.warning("Harap masukkan instruksi perbaikan terlebih dahulu.")
                else:
                    with st.spinner("AI sedang merevisi silabus berdasarkan instruksi Anda..."):
                        fix_prompt = f"""
                        Anda adalah pakar kurikulum madrasah/pesantren.
                        Berikut adalah data silabus saat ini dalam format JSON:
                        {json.dumps(data, ensure_ascii=False, indent=2)}

                        User memberikan instruksi/perintah perbaikan berikut:
                        "{correction_prompt}"

                        Terapkan perbaikan tersebut secara cermat dan proporsional. 
                        Keluarkan respons HANYA dalam bentuk JSON valid dengan struktur yang sama persis seperti sebelumnya tanpa teks lain:
                        {{
                            "nama_kitab": "{data.get('nama_kitab', '')}",
                            "fan_ilmu": "{data.get('fan_ilmu', '')}",
                            "tingkat_kelas": "{data.get('tingkat_kelas', '')}",
                            "total_pertemuan": "{data.get('total_pertemuan', '')}",
                            "silabus": [
                                {{
                                    "ptm": 1,
                                    "bab": "...",
                                    "halaman": "...",
                                    "alokasi_waktu": "...",
                                    "metode": "...",
                                    "capaian": "...",
                                    "indikator": "...",
                                    "evaluasi": "..."
                                }}
                            ]
                        }}
                        """
                        try:
                            fix_response = client.models.generate_content(
                                model="gemini-3.6-flash",
                                contents=fix_prompt,
                                config=types.GenerateContentConfig(max_output_tokens=65536, temperature=0.7)
                            )
                            raw_fix = fix_response.text.strip()
                            if "```json" in raw_fix:
                                raw_fix = raw_fix.split("```json")[1].split("```")[0].strip()
                            elif "```" in raw_fix:
                                raw_fix = raw_fix.split("```")[1].split("```")[0].strip()

                            start_idx = raw_fix.find("{")
                            end_idx = raw_fix.rfind("}")
                            if start_idx != -1 and end_idx != -1:
                                raw_fix = raw_fix[start_idx : end_idx + 1]

                            updated_data = json.loads(raw_fix)
                            st.session_state["result_data"] = updated_data
                            st.session_state["rpp_cache"] = {}
                            st.success("✅ Silabus berhasil diperbarui sesuai instruksi Anda!")
                            st.rerun()
                        except Exception as fix_err:
                            st.error(f"Gagal merevisi silabus: {fix_err}")
            
            st.markdown("---")
            silabus_docx_buffer = create_silabus_docx(data)
            st.download_button(
                label="📥 Download Dokumen Silabus (.docx)",
                data=silabus_docx_buffer,
                file_name=f"Silabus_{data.get('nama_kitab','Kitab').replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

    with tab_jurnal:
        st.markdown("### Preview Jurnal Pembelajaran KBM (Excel)")
        st.caption("Tabel jurnal harian pengajar disesuaikan dengan draf Excel Anda.")
        
        silabus_list = data.get("silabus", [])
        if silabus_list:
            df_jurnal = pd.DataFrame([{
                "MATERI POKOK": item.get("bab", ""),
                "ALOKASI WAKTU": item.get("alokasi_waktu", ""),
                "Tanggal": "......./....../2026",
                "NAMA PENGAJAR/PENGGANTI": "",
                "CATATAN MATERI (Diisi jika diperlukan)": "",
                "TTD": "",
                "Keterangan": ""
            } for item in silabus_list])
            
            st.dataframe(df_jurnal, use_container_width=True, hide_index=True)
            
            jurnal_excel_buffer = create_jurnal_excel(data)
            st.download_button(
                label="📥 Download Jurnal Mengajar (.xlsx)",
                data=jurnal_excel_buffer,
                file_name=f"Jurnal_Mengajar_{data.get('nama_kitab','Kitab').replace(' ', '_')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            
    with tab_rpp:
        silabus_items = data.get("silabus", [])
        if silabus_items:
            ptm_options = [f"Pertemuan Ke-{item.get('ptm')} ({item.get('bab', 'Materi')})" for item in silabus_items]
            selected_ptm_label = st.selectbox("Pilih Pertemuan untuk Generate RPP:", ptm_options)
            
            selected_idx = ptm_options.index(selected_ptm_label)
            selected_item = silabus_items[selected_idx]
            ptm_num = selected_item.get("ptm")
            
            if st.button(f"✨ Buat RPP untuk Pertemuan Ke-{ptm_num}", type="secondary"):
                with st.spinner(f"AI sedang merancang RPP khusus untuk Pertemuan Ke-{ptm_num}..."):
                    rpp_prompt = f"""
                    Anda adalah pakar kurikulum pesantren. Buatkan SATU Rencana Pelaksanaan Pembelajaran (RPP) yang sangat detail untuk Pertemuan ke-{ptm_num} berdasarkan data silabus berikut:
                    - Bab/Materi: {selected_item.get('bab')}
                    - Halaman: {selected_item.get('halaman')}
                    - Capaian: {selected_item.get('capaian')}
                    - Indikator: {selected_item.get('indikator')}
                    - Alokasi Waktu: {selected_item.get('alokasi_waktu')}
                    
                    Gunakan teks sumber kitab jika relevan. Berikan keluaran HANYA dalam format JSON valid berikut tanpa teks lain:
                    {{
                        "ptm": {ptm_num},
                        "sekolah": "MDT / Pesantren Rufidz Tahfidz & Diniyah Indonesia",
                        "materi_pokok": "{selected_item.get('bab')}",
                        "alokasi_waktu": "{selected_item.get('alokasi_waktu')}",
                        "tujuan_pembelajaran": [
                            "Tujuan 1...",
                            "Tujuan 2..."
                        ],
                        "langkah_pembelajaran": {{
                            "pendahuluan": [
                                "Guru membuka pembelajaran dengan salam dan doa.",
                                "Guru menyampaikan apersepsi."
                            ],
                            "inti": [
                                {{"aspek": "Kegiatan Literasi", "kegiatan": "Santri menyimak..."}},
                                {{"aspek": "Critical Thinking", "kegiatan": "Santri menganalisis..."}},
                                {{"aspek": "Collaboration", "kegiatan": "Diskusi kelompok..."}},
                                {{"aspek": "Communication", "kegiatan": "Presentasi hasil..."}},
                                {{"aspek": "Creativity", "kegiatan": "Membuat kesimpulan..."}}
                            ],
                            "penutup": [
                                "Guru bersama santri menyimpulkan materi.",
                                "Doa penutup majelis."
                            ]
                        }},
                        "penilaian": {{
                            "sikap": "Observasi keaktifan dan kedisiplinan santri.",
                            "pengetahuan": "Tes lisan / Tanya jawab wazan/kaidah.",
                            "keterampilan": "Praktik membaca teks dan tashrif."
                        }}
                    }}
                    """
                    try:
                        rpp_response = client.models.generate_content(
                            model="gemini-3.6-flash",
                            contents=rpp_prompt,
                            config=types.GenerateContentConfig(max_output_tokens=65536, temperature=0.7)
                        )
                        raw_rpp = rpp_response.text.strip()
                        if raw_rpp.startswith("```"):
                            lines = raw_rpp.splitlines()
                            if lines[0].startswith("```"): lines = lines[1:]
                            if lines and lines[-1].startswith("```"): lines = lines[:-1]
                            raw_rpp = "\n".join(lines).strip()
                        
                        rpp_data = json.loads(raw_rpp)
                        st.session_state["rpp_cache"][ptm_num] = rpp_data
                        st.success(f"✅ RPP Pertemuan ke-{ptm_num} berhasil dibuat!")
                    except Exception as e:
                        st.error(f"Gagal generate RPP: {str(e)}")
            
            if ptm_num in st.session_state["rpp_cache"]:
                rpp = st.session_state["rpp_cache"][ptm_num]
                
                st.markdown("---")
                st.markdown("### THE RUFIDZ INDONESIA")
                st.markdown("#### RENCANA PELAKSANAAN PEMBELAJARAN (RPP)")
                st.markdown(f"**Sekolah**: {rpp.get('sekolah', '-')}")
                st.markdown(f"**Mata Pelajaran**: {data.get('fan_ilmu', '')} ({data.get('nama_kitab', '')})")
                st.markdown(f"**Kelas / Tahap**: {data.get('tingkat_kelas', '-')}")
                st.markdown(f"**Materi Pokok**: {rpp.get('materi_pokok', '-')}")
                st.markdown(f"**Alokasi Waktu**: {rpp.get('alokasi_waktu', '-')}")
                
                st.markdown("#### A. Tujuan Pembelajaran")
                for t in rpp.get("tujuan_pembelajaran", []):
                    st.markdown(f"• {t}")
                    
                st.markdown("#### B. Langkah-Langkah Pembelajaran")
                st.markdown("**1. Kegiatan Pendahuluan**")
                for p_item in rpp.get("langkah_pembelajaran", {}).get("pendahuluan", []):
                    st.markdown(f"• {p_item}")
                    
                st.markdown("**2. Kegiatan Inti**")
                df_inti = pd.DataFrame(rpp.get("langkah_pembelajaran", {}).get("inti", []))
                if not df_inti.empty:
                    if "aspek" in df_inti.columns and "kegiatan" in df_inti.columns:
                        df_inti.rename(columns={"aspek": "Aspek", "kegiatan": "Kegiatan Pembelajaran"}, inplace=True)
                    st.table(df_inti)
                    
                st.markdown("**3. Kegiatan Penutup**")
                for p_item in rpp.get("langkah_pembelajaran", {}).get("penutup", []):
                    st.markdown(f"• {p_item}")
                    
                st.markdown("#### C. Penilaian Hasil Pembelajaran")
                pen = rpp.get("penilaian", {})
                st.write(f"• **Penilaian Sikap**: {pen.get('sikap')}")
                st.write(f"• **Penilaian Pengetahuan**: {pen.get('pengetahuan')}")
                st.write(f"• **Penilaian Keterampilan**: {pen.get('keterampilan')}")
                
                docx_buffer = create_single_rpp_docx(data, rpp)
                st.download_button(
                    label=f"📥 Download RPP Pertemuan ke-{ptm_num} (.docx)",
                    data=docx_buffer,
                    file_name=f"RPP_Pertemuan_{ptm_num}_{data.get('nama_kitab','').replace(' ', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
